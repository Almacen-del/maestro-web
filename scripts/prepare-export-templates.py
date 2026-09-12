"""Sanitize data cells in copies, preserving native OOXML parts and cell styles.
openpyxl is used only to inspect ranges; authoring patches the ZIP XML directly.
"""
from pathlib import Path
import re,json,zipfile,hashlib
import openpyxl
from openpyxl.utils.cell import range_boundaries,get_column_letter
ROOT=Path('C:/Users/Almacen/Downloads')
OUT=Path('public/export-templates');OUT.mkdir(parents=True,exist_ok=True)
FILES={'daily':'FORMATO PROCESO DIARIOS 2026.xlsx','applications':'PV-F-001 CONTROL DE APLICACIONES.xlsx','grafts':'PV-F 008 HOJA DE VERIFICACION  DE TRAZABILIDAD DE INJERTACIÓN (1).xlsx','buds':'PV-F-007 HOJA DE VERIFICACION  DE TRAZABILIDAD DE YEMAS.xlsx','monitoring':'PV-F-010 HOJA DE MONITOREO DE VIVERO.xlsx','inventory':'INVENTARIO JUNIO 2026 (2).xlsx'}
def cells(ranges):
    result=set()
    for value in ranges:
        a,b,c,d=range_boundaries(value)
        result.update(f'{get_column_letter(col)}{row}' for row in range(b,d+1) for col in range(a,c+1))
    return result
manifest={}
for kind,name in FILES.items():
    print('preparing',kind,flush=True)
    source=ROOT/name;original_hash=hashlib.sha256(source.read_bytes()).hexdigest()
    wb=openpyxl.load_workbook(source);z=zipfile.ZipFile(source);parts={n:z.read(n) for n in z.namelist()};sheets=[]
    for index,s in enumerate(wb.worksheets):
        if 'CAMBIO' in s.title.upper():continue
        ranges=[];meta={'name':s.title,'path':f'xl/worksheets/sheet{index+1}.xml'}
        if kind in ['daily','applications']:ranges=[f'A5:{get_column_letter(s.max_column)}{s.max_row}'];meta.update(start=5,capacity=950 if kind=='daily' else 999)
        elif kind=='grafts':ranges=['D5','L5','C8','C12','G12','B16:M37','D38','B40'];meta.update(capacity=22)
        elif kind=='buds':
            ranges=['B4','E4','C53','A55']
            for r in [11,18,25,32,39,46]:ranges += [f'B{r}',f'D{r}',f'F{r}',f'B{r+1}',f'D{r+1}',f'B{r+4}:F{r+6}']
            meta.update(capacity=6)
        elif kind=='monitoring':
            ranges=['C5','C6','C55','A57','B14','B21','B30','B38','B44','B54','C25','G24','J25','C31','C46']
            for lo,hi in [(11,13),(19,20),(28,29),(34,37),(41,43),(49,53)]:ranges.append(f'C{lo}:J{hi}')
            # Existing quantity labels are preserved; historical formulas in input cells are cleared.
        else:
            start=9 if index==0 else 8;col=1 if index==0 else 2
            total=next(c.row for row in s for c in row if c.column==col and str(c.value).strip().lower()=='total')
            ranges=[f'{get_column_letter(col)}{start}:{get_column_letter(col+7)}{total-1}',f'{get_column_letter(col+3)}{total}:{get_column_letter(col+6)}{total}']
            ranges+=['B6','D6','F6','H6','D87','A89'] if index==0 else ['C5','E5','G5','I5']
            for row in s:
                for c in row:
                    if isinstance(c.value,str) and c.value.startswith('FIRMA DE DIRECTO'):ranges.append(f'{get_column_letter(col+3)}{c.row}')
                    if isinstance(c.value,str) and c.value.strip()=='OBSERVACIONES' and c.row>total:ranges.append(f'{get_column_letter(col)}{c.row+1}')
            meta.update(start=start,column=col,totalRow=total,capacity=total-start)
        allowed=cells(ranges);xml=parts[meta['path']].decode('utf-8')
        def clean(m):
            ref=re.search(r'\br="([A-Z]+\d+)"',m[0])
            if not ref or ref[1] not in allowed:return m[0]
            if kind=='inventory' and ref[1].endswith(str(meta['totalRow'])) and '<f' in m[0]:
                return re.sub(r'<v>.*?</v>','<v>0</v>',m[0],flags=re.S)
            tag=m[0].split('>')[0].rstrip('/')
            tag=re.sub(r'\s+t="[^"]*"','',tag)
            return tag+'/>'
        xml=re.sub(r'<c\b[^>]*?(?:/>|>.*?</c>)',clean,xml,flags=re.S)
        parts[meta['path']]=xml.encode();meta['allowed']=sorted(allowed);sheets.append(meta)
    # Empty unused shared strings so removed history is not retained inside the ZIP.
    if 'xl/sharedStrings.xml' in parts:
        used=set()
        for path,data in parts.items():
            if re.match(r'xl/worksheets/sheet\d+\.xml$',path):
                for tag in re.findall(r'<c\b[^>]*?(?:/>|>.*?</c>)',data.decode(),re.S):
                    if re.search(r'\bt="s"',tag):
                        v=re.search(r'<v>(\d+)</v>',tag)
                        if v:used.add(int(v[1]))
        counter=[-1]
        def scrub(m):
            counter[0]+=1
            return m[0] if counter[0] in used else '<si><t></t></si>'
        parts['xl/sharedStrings.xml']=re.sub(r'<si\b[^>]*>.*?</si>',scrub,parts['xl/sharedStrings.xml'].decode(),flags=re.S).encode()
    with zipfile.ZipFile(OUT/f'{kind}.xlsx','w',zipfile.ZIP_DEFLATED) as dest:
        for path,data in parts.items():dest.writestr(path,data)
    assert hashlib.sha256(source.read_bytes()).hexdigest()==original_hash
    manifest[kind]={'sheets':sheets,'sourceHash':original_hash}
    print(kind,'sanitized',len(sheets),'sheets; source unchanged')
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False),encoding='utf-8')
